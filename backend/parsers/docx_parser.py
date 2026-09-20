import io
from pathlib import Path
from typing import Union
import docx
from .cleaner import clean_text

def extract_text_from_docx(source: Union[str, Path, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a DOCX file path or byte stream using python-docx.
    Extracts from paragraphs and tables.
    Returns cleaned text.
    """
    try:
        if isinstance(source, (str, Path)):
            doc = docx.Document(str(source))
        elif isinstance(source, bytes):
            doc = docx.Document(io.BytesIO(source))
        else:
            doc = docx.Document(source)
            
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text (frequent in resume contact/education headers)
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Avoid duplicated adjacent identical cells (merged cells)
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    table_texts.append(" | ".join(unique_cells))
                    
        combined = "\n".join(paragraphs)
        if table_texts:
            combined += "\n\n" + "\n".join(table_texts)
            
        return clean_text(combined)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")
